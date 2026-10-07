"""Generate the DxPS installation and running guide for Windows, Ubuntu and macOS.

One content model is rendered twice: dxps/INSTALL.md (ships with the source) and
DxPS_Installation_and_Running_Guide.docx (converted to PDF with Word).
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from build_dxps_tobe_docx import GREY, NAVY, add_toc, bullets, code, h, page_break, para, table
from build_dxps_testcases_docx import header_footer, kpi_table

HERE = Path(__file__).parent
OUT_DOCX = HERE.parent / "DxPS_Installation_and_Running_Guide.docx"
OUT_MD = HERE.parent / "dxps" / "INSTALL.md"

ORDER_JSON = """{
  "@type": "ServiceOrder",
  "externalId": "DEMO-0001",
  "channel": [{"name": "SELFCARE"}],
  "serviceOrderItem": [{
    "id": "1", "action": "add",
    "service": {
      "serviceSpecification": {"id": "CreateSubscriber5G"},
      "serviceCharacteristic": [
        {"name": "supi",   "value": "imsi-416770000004321"},
        {"name": "msisdn", "value": "962790004321"},
        {"name": "plan",   "value": "5G-100GB"}
      ]
    }
  }]
}"""

# Content model: (kind, ...) blocks.
#   ("h1"|"h2"|"h3", text) ("p", text) ("bullets", [..]) ("steps", [(title, text, code_lang, code)])
#   ("code", lang, text) ("table", header, rows) ("note", text)
C = []
add = C.append

add(("h1", "1. What you are installing"))
add(("p", "DxPS is a multi-tenant real-time provisioning system for MNOs and MVNOs. It takes TM Forum TMF641 service "
          "orders, plans them into network element (NE) tasks, executes them with sagas over Kafka and reports progress "
          "with TMF688 events. The package runs the complete system on one machine: PostgreSQL 18, Apache Kafka 4, "
          "six DxPS services, NE simulators for every southbound interface, and the Mosaic live dashboard."))
add(("table", ["Component", "Address (loopback only)", "Purpose"], [
    ["Mosaic dashboard", "http://127.0.0.1:8088", "Live view: tenants, lanes, topics, orders, NEs, events, tests; load generator and attack probes"],
    ["Gateway (TMF641/TMF688 hub)", "https://localhost:8443", "Northbound API, JWT auth, per-tenant rate limits"],
    ["Orchestrator", "127.0.0.1:9201 (stats)", "Plans orders, runs sagas, dependencies and per-subscriber sequencing"],
    ["Adapter", "127.0.0.1:9202 (stats)", "Southbound drivers: 5G SBA, IMS, NETCONF/USP, eSIM, BSS/OCS, NPDB, CAMARA"],
    ["Event hub", "127.0.0.1:9203 (stats)", "Signed TMF688 webhooks with retries"],
    ["NE simulators", "9101-9108 (mTLS), admin 9199, webhook sink 9109", "NRF, UDR, IMS, RESTCONF/USP, SM-DP+, OCS/BSS, NPDB, NEF/CAMARA"],
    ["PostgreSQL 18", "localhost:5433 (TLS 1.3, SCRAM)", "Orders, plans, tasks, outbox, audit; row-level security per tenant"],
    ["Apache Kafka 4 (KRaft)", "127.0.0.1:9092", "41 topics: business commands, NE tasks, retry ladder, events, DLQ"],
]))
add(("note", "Every listener binds to the loopback interface. Nothing is exposed to the network."))

add(("h1", "2. Requirements"))
add(("table", ["Item", "Minimum", "Recommended"], [
    ["CPU", "4 cores (x86-64 or ARM64)", "8 cores"],
    ["Memory", "8 GB", "16 GB"],
    ["Disk", "6 GB free (tools, data, logs)", "15 GB"],
    ["Windows", "Windows 10 22H2 / Windows 11, x64, PowerShell 5.1+, winget", "Windows 11"],
    ["Ubuntu", "22.04 LTS or 24.04 LTS, a normal user with sudo", "24.04 LTS"],
    ["macOS", "macOS 13 Ventura or later, Homebrew, Xcode Command Line Tools", "macOS 14+ on Apple silicon"],
    ["Network", "HTTPS to go.dev/dl.google.com, downloads.apache.org (or archive.apache.org), "
                "apt.postgresql.org / Homebrew / get.enterprisedb.com, proxy.golang.org", ""],
    ["Free ports", "5433, 8088, 8443, 9092, 9093, 9101-9109, 9199, 9201-9203", ""],
]))
add(("p", "The installer scripts fetch these versions (all can be overridden, see section 8):"))
add(("table", ["Software", "Windows", "Ubuntu", "macOS"], [
    ["Go 1.26.6", "winget GoLang.Go", "official tarball, SHA-256 verified, into the runtime", "Homebrew go"],
    ["Java 21 (Kafka needs 17+)", "winget Eclipse Temurin 21 JRE", "apt openjdk-21-jre-headless", "Homebrew openjdk@21"],
    ["PostgreSQL 18", "EDB binaries zip, into the runtime", "apt.postgresql.org postgresql-18", "Homebrew postgresql@18"],
    ["Apache Kafka 4.3.1", "Apache release, SHA-512 verified", "Apache release, SHA-512 verified", "Apache release, SHA-512 verified"],
]))

add(("h1", "3. Package contents"))
add(("p", "The release (GitHub release v<version>, or the DxPS_Delivery_<date> folder) contains:"))
add(("code", "text", """dxps-<version>-windows.zip        Windows bundle (source, scripts, manuals)
dxps-<version>-ubuntu.tar.gz      Ubuntu / macOS bundle (same content; scripts keep their execute bit)
dxps-<version>-evidence.zip       test evidence: tests.json, coverage, probes, govulncheck, gosec, payload traces
SHA256SUMS.txt                    checksums of the files above

dxps/ (inside each bundle)
  RELEASE.txt  start here: contents and quick start for the platform
  cmd/         gateway orchestrator adapter eventhub netsim dashboard dxpsctl
  internal/    24 packages (catalog, planner, saga, bus, store, ...) with their tests
  scripts/     *.ps1 for Windows, *.sh for Ubuntu and macOS
  docs/manuals/  HLD+LLD, Test Cases and Results, Components and Data Guide, this guide (docx + pdf)
  INSTALL.md   this guide in Markdown"""))
add(("p", "All scripts keep their state in one runtime directory, by default %USERPROFILE%\\dxps-runtime on Windows "
          "and ~/dxps-runtime on Ubuntu and macOS (override with DXPS_RUNTIME). The source tree is never modified."))
add(("table", ["Script (Windows .ps1 / Ubuntu+macOS .sh)", "What it does", "When"], [
    ["install-deps-windows / -ubuntu / -macos", "Installs Go, Java, PostgreSQL 18 and Kafka 4", "once"],
    ["infra-setup", "Generates secrets and the dev PKI, creates the PostgreSQL cluster, roles and database, formats Kafka", "once"],
    ["run-all", "Starts PostgreSQL and Kafka, builds all services, migrates, creates topics, seeds reference data, starts the stack", "every start"],
    ["run-all -NoBuild / --no-build", "Same without rebuilding", "fast restart"],
    ["status", "Shows which components are listening", "any time"],
    ["test-all", "Runs all unit, integration and E2E tests with coverage; publishes results to the Mosaic Tests tile", "after changes"],
    ["sec-scan", "govulncheck + gosec static security scans", "before delivery"],
    ["infra-stop", "Stops services, Kafka and PostgreSQL", "end of day"],
    ["infra-start", "Starts only PostgreSQL and Kafka", "for development"],
]))

# ---------------------------------------------------------------- Windows
add(("h1", "4. Windows 10/11 - step by step"))
add(("steps", [
    ("Unpack the source", "Extract dxps-<version>-windows.zip, for example to C:\\ (it creates C:\\dxps). Open Windows PowerShell in that folder "
     "(the folder that contains go.mod) and allow local scripts for this window only.", "powershell",
     "cd C:\\dxps\nSet-ExecutionPolicy -Scope Process -ExecutionPolicy Bypass"),
    ("Install the prerequisites", "Installs Go and Java with winget if missing, and downloads PostgreSQL 18 binaries and "
     "Kafka 4 into the runtime directory. When winget installs Go or Java, close PowerShell, open a new window and "
     "repeat the command so PATH is refreshed.", "powershell", ".\\scripts\\install-deps-windows.ps1"),
    ("Create the runtime", "Generates secrets.env (random passwords and keys, never printed), the dev CA and service "
     "certificates, initialises PostgreSQL on port 5433 with TLS 1.3 and SCRAM, creates the roles dxps_owner, dxps_app "
     "(row-level security enforced) and dxps_ops, and formats Kafka in KRaft mode.", "powershell", ".\\scripts\\infra-setup.ps1"),
    ("Build and start everything", "Takes 1-3 minutes the first time (Go module download and build).", "powershell",
     ".\\scripts\\run-all.ps1"),
    ("Open the Mosaic", "Browse to http://127.0.0.1:8088. Press Start load to send a realistic order mix, Run probes to attack "
     "the live system with 24 security probes, and use the chaos slider to inject NE faults.", None, None),
    ("Check and test", "", "powershell", ".\\scripts\\status.ps1\n.\\scripts\\test-all.ps1      # ~2 minutes, results appear in the Tests tile\n.\\scripts\\sec-scan.ps1"),
    ("Stop and restart", "", "powershell", ".\\scripts\\infra-stop.ps1\n.\\scripts\\run-all.ps1 -NoBuild"),
]))
add(("note", "Windows-specific: Kafka on Windows cannot rename open log segments (KAFKA-1194), so the dev broker runs with "
             "log compaction and retention deletion disabled. Do not delete Kafka topics on Windows; it takes the broker "
             "down. Linux and macOS use Kafka's defaults."))

# ---------------------------------------------------------------- Ubuntu
add(("h1", "5. Ubuntu 22.04 / 24.04 - step by step"))
add(("steps", [
    ("Unpack the source", "Use a normal user with sudo rights, not root: PostgreSQL refuses to run as root.", "bash",
     "mkdir -p ~/dxps && tar -xzf dxps-<version>-ubuntu.tar.gz -C ~/dxps --strip-components=1\ncd ~/dxps"),
    ("Install the prerequisites", "Adds the official PostgreSQL apt repository and installs postgresql-18 without creating "
     "the system cluster, installs OpenJDK 21, installs Go 1.26.6 into the runtime if the system Go is older, and "
     "downloads Kafka 4. Asks for your sudo password.", "bash", "bash scripts/install-deps-ubuntu.sh"),
    ("Create the runtime", "Same as on Windows: secrets, PKI, PostgreSQL cluster on 5433 with TLS 1.3 + SCRAM, roles, "
     "database, Kafka KRaft storage.", "bash", "bash scripts/infra-setup.sh"),
    ("Build and start everything", "", "bash", "bash scripts/run-all.sh"),
    ("Open the Mosaic", "Browse to http://127.0.0.1:8088. On a remote server, tunnel it instead of exposing it: "
     "ssh -L 8088:127.0.0.1:8088 user@server, then open http://127.0.0.1:8088 locally.", None, None),
    ("Check and test", "", "bash", "bash scripts/status.sh\nbash scripts/test-all.sh\nbash scripts/sec-scan.sh"),
    ("Stop and restart", "", "bash", "bash scripts/infra-stop.sh\nbash scripts/run-all.sh --no-build"),
]))

# ---------------------------------------------------------------- macOS
add(("h1", "6. macOS 13+ - step by step"))
add(("steps", [
    ("Install Homebrew and the command line tools", "Skip if you already have them.", "bash",
     "xcode-select --install\n/bin/bash -c \"$(curl -fsSL https://raw.githubusercontent.com/Homebrew/install/HEAD/install.sh)\""),
    ("Unpack the source", "Use the Ubuntu bundle; it is the same for macOS.", "bash",
     "mkdir -p ~/dxps && tar -xzf dxps-<version>-ubuntu.tar.gz -C ~/dxps --strip-components=1\ncd ~/dxps"),
    ("Install the prerequisites", "Installs go, postgresql@18 and openjdk@21 with Homebrew (no brew services are started) and "
     "downloads Kafka 4 into the runtime.", "bash", "bash scripts/install-deps-macos.sh"),
    ("Create the runtime", "", "bash", "bash scripts/infra-setup.sh"),
    ("Build and start everything", "macOS may ask once whether the services may accept incoming connections. They only "
     "listen on loopback; Allow or Deny both work for local use.", "bash", "bash scripts/run-all.sh"),
    ("Open the Mosaic", "Browse to http://127.0.0.1:8088.", None, None),
    ("Check and test", "", "bash", "bash scripts/status.sh\nbash scripts/test-all.sh\nbash scripts/sec-scan.sh"),
    ("Stop and restart", "", "bash", "bash scripts/infra-stop.sh\nbash scripts/run-all.sh --no-build"),
]))

# ---------------------------------------------------------------- using it
add(("h1", "7. Using the running system"))
add(("h2", "7.1 Submit a TMF641 order"))
add(("p", "Save this order as order.json. Subscriber numbers ending in 997, 998 or 999 trigger simulator faults "
          "(transient, permanent rejection, unavailable), so this example avoids them."))
add(("code", "json", ORDER_JSON))
add(("p", "Easiest on every platform: dxpsctl api mints a short-lived token for the tenant internally (never printed) "
          "and calls the gateway over TLS with the dev CA."))
add(("code", "bash", "# Ubuntu / macOS\n~/dxps-runtime/bin/dxpsctl api -tenant mvno-alpha -idem demo-0001 -body \"$(cat order.json)\" POST /serviceOrder\n"
                     "~/dxps-runtime/bin/dxpsctl api -tenant mvno-alpha GET \"/serviceOrder?limit=5\""))
add(("code", "powershell", "# Windows\n$ctl = \"$env:USERPROFILE\\dxps-runtime\\bin\\dxpsctl.exe\"\n"
                           "& $ctl api -tenant mvno-alpha -idem demo-0001 -body (Get-Content -Raw order.json) POST /serviceOrder\n"
                           "& $ctl api -tenant mvno-alpha GET \"/serviceOrder?limit=5\""))
add(("p", "With curl (Ubuntu/macOS). The token is a credential: keep it in a variable, do not paste it into tickets or chat."))
add(("code", "bash", "RT=~/dxps-runtime\nTOKEN=$($RT/bin/dxpsctl token -tenant mvno-alpha -ttl 10m)\n"
                     "curl --cacert $RT/pki/ca.crt -H \"Authorization: Bearer $TOKEN\" -H 'Content-Type: application/json' \\\n"
                     "     -H 'Idempotency-Key: demo-0002' --data @order.json \\\n"
                     "     https://localhost:8443/tmf-api/serviceOrdering/v5/serviceOrder\nunset TOKEN"))
add(("p", "The gateway answers 201 with the order id and state acknowledged. Sending the same Idempotency-Key again "
          "returns 200 with Idempotent-Replayed: true. Follow the order in the Mosaic Orders and Events tiles."))
add(("h2", "7.2 Tenants"))
add(("table", ["Tenant", "Type", "Notes"], [
    ["host-mno", "Host MNO", "owns the shared core network"],
    ["mvno-alpha", "Full MVNO", "own OCS and a tenant-specific CreateSubscriber5G variant"],
    ["mvno-beta", "Light MVNO", "uses the host's shared NEs"],
    ["ent-acme", "Enterprise", "enterprise customer of the host MNO"],
    ["mvno-gamma", "Light MVNO", "suspended: every call is refused (403 DXPS-1004)"],
]))
add(("h2", "7.3 Logs, data and results"))
add(("table", ["Path (under the runtime directory)", "Content"], [
    ["logs/svc/<service>.log", "JSON logs of each DxPS service (no secrets)"],
    ["logs/postgres.log, logs/kafka/server.log", "PostgreSQL and Kafka logs"],
    ["results/tests.json, coverage.out, tests-raw.jsonl", "latest test-all run"],
    ["results/govulncheck.txt, gosec.json", "latest sec-scan run"],
    ["secrets.env (mode 0600)", "generated passwords, DSNs and keys - never share or commit"],
    ["pki/", "dev CA and certificates (development only)"],
    ["data/pg, data/kafka", "PostgreSQL cluster and Kafka log directories"],
]))
add(("h2", "7.4 Regenerate the test-case document"))
add(("p", "After test-all, the Word report can be rebuilt from the new results with Python 3.10+, python-docx and "
          "matplotlib (design/build_dxps_testcases_docx.py in the delivery workspace). On Windows, Word exports the PDF; "
          "on Ubuntu/macOS use LibreOffice: soffice --headless --convert-to pdf DxPS_Test_Cases_and_Results.docx."))

# ---------------------------------------------------------------- config
add(("h1", "8. Configuration reference"))
add(("table", ["Variable", "Default", "Meaning"], [
    ["DXPS_RUNTIME", "~/dxps-runtime (Windows: %USERPROFILE%\\dxps-runtime)", "runtime directory for tools, data, logs, secrets"],
    ["DXPS_PGBIN", "Ubuntu /usr/lib/postgresql/18/bin; macOS $(brew --prefix postgresql@18)/bin", "PostgreSQL binaries (Ubuntu/macOS scripts)"],
    ["DXPS_PG_MAJOR", "18", "PostgreSQL major version for the Ubuntu/macOS installers"],
    ["DXPS_KAFKA_VERSION", "4.3.1", "Kafka release downloaded by the installers (Windows: -KafkaVersion)"],
    ["DXPS_GO_VERSION", "1.26.6", "Go version installed on Ubuntu when the system Go is older"],
    ["KAFKA_HEAP_OPTS", "-Xms512m -Xmx1g", "Kafka JVM heap"],
    ["DXPS_NETSIM_HOST", "localhost", "host name the seed uses for simulator endpoints"],
    ["HTTPS_PROXY / GOPROXY", "-", "corporate proxy for downloads and Go modules"],
]))

# ---------------------------------------------------------------- troubleshooting
add(("h1", "9. Troubleshooting"))
add(("table", ["Symptom", "Cause and fix"], [
    ["'port NNNN did not open'", "Another program uses the port, or the component failed: run status, then read the "
                                 "component's log (section 7.3). On Linux: ss -ltnp | grep NNNN; macOS: lsof -iTCP:NNNN -sTCP:LISTEN."],
    ["Kafka does not start", "java -version must be 17 or later. Read logs/kafka/server.log. On macOS check that "
                             "openjdk@21 is installed (the scripts set JAVA_HOME automatically)."],
    ["Windows: Kafka stops with 'all log dirs have failed'", "A topic was deleted or segments were renamed. Stop the stack, delete "
                                                             "only the partition folders named in server.log, start again."],
    ["'initdb: cannot be run as root'", "Run the Ubuntu/macOS scripts as a normal user with sudo rights."],
    ["Ubuntu: 'could not create lock file /var/run/postgresql/...'", "The cluster tried to use the system socket directory. "
                                                                    "Re-run infra-setup: it sets unix_socket_directories = '' "
                                                                    "(DxPS uses TLS over TCP only)."],
    ["PostgreSQL 'could not load server certificate'", "Certificate permissions changed. The files in data/pg must be owned "
                                                       "by your user with mode 0600; re-copy them from pki/."],
    ["macOS: postgresql@18 not found", "brew update && brew install postgresql@18, or set DXPS_PGBIN to another PostgreSQL 18 bin directory."],
    ["Go build downloads fail", "Set HTTPS_PROXY (and GOPROXY if your company mirrors modules)."],
    ["Dashboard shows a service as down", "Restart with run-all; logs/svc/<service>.out shows start-up errors."],
    ["Start from scratch", "Stop the stack, then remove the runtime directory (this deletes all data and secrets) "
                           "and rerun infra-setup."],
]))

add(("h1", "10. Security notes"))
add(("bullets", [
    "Secrets are generated locally with crypto/rand, stored only in secrets.env (mode 0600) and never printed by the scripts.",
    "PostgreSQL accepts only TLS 1.3 + SCRAM-SHA-256 on loopback; the application role is subject to row-level security.",
    "All NE simulator links use mutual TLS; the gateway uses TLS with the dev CA; webhooks are HMAC-signed.",
    "The dev CA, PLAINTEXT Kafka on loopback and the single-node layout are for development and demonstration. "
    "Production uses the HLD/LLD deployment (Kafka SASL_SSL, managed PKI/HSM, Kubernetes).",
]))
add(("h1", "11. Uninstall"))
add(("bullets", [
    "Stop the stack with infra-stop, then delete the runtime directory and the source folder.",
    "Windows: winget uninstall GoLang.Go / EclipseAdoptium.Temurin.21.JRE (only if they were installed for DxPS).",
    "Ubuntu: sudo apt remove postgresql-18 openjdk-21-jre-headless; sudo rm /etc/apt/sources.list.d/pgdg.list.",
    "macOS: brew uninstall postgresql@18 openjdk@21 go.",
]))


# --------------------------------------------------------------------------- renderers
def to_markdown():
    out = ["# DxPS - Installation and Running Guide (Windows, Ubuntu, macOS)", ""]
    for b in C:
        k = b[0]
        if k in ("h1", "h2", "h3"):
            out += [("#" * (int(k[1]) + 1)) + " " + b[1], ""]
        elif k == "p":
            out += [b[1], ""]
        elif k == "note":
            out += ["> **Note:** " + b[1], ""]
        elif k == "bullets":
            out += ["- " + x for x in b[1]] + [""]
        elif k == "code":
            out += ["```" + ("" if b[1] == "text" else b[1]), b[2], "```", ""]
        elif k == "table":
            hdr, rows = b[1], b[2]
            out += ["| " + " | ".join(hdr) + " |", "|" + "|".join("---" for _ in hdr) + "|"]
            out += ["| " + " | ".join(str(c).replace("|", "\\|") for c in r) + " |" for r in rows] + [""]
        elif k == "steps":
            for i, (title, text, lang, src) in enumerate(b[1], 1):
                out += [f"**Step {i} - {title}.** {text}".rstrip(), ""]
                if src:
                    out += ["```" + lang, src, "```", ""]
    return "\n".join(out).rstrip() + "\n"


def to_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.left_margin = sec.right_margin = Inches(0.75)
    sec.top_margin = sec.bottom_margin = Inches(0.75)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    header_footer(doc, "DxPS Real-time Provisioning System - Installation and Running Guide")
    for _ in range(4):
        doc.add_paragraph()
    para(doc, "DxPS", bold=True, size=30, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    para(doc, "Real-time Provisioning System", bold=True, size=20, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "Installation and Running Guide", bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "Windows 10/11  -  Ubuntu 22.04/24.04  -  macOS 13+", size=12, italic=True, color=GREY,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)
    kpi_table(doc, [("steps to a running stack", "4", None), ("platforms", "3", None),
                    ("services + simulators", "6 + 9", None), ("dashboard", "127.0.0.1:8088", None)])
    table(doc, ["Item", "Value"], [
        ["Document", "DXPS-INSTALL-GUIDE"], ["Version", "1.1"], ["Date", "07 October 2026"],
        ["Applies to", "DxPS release 1.0.0; Go 1.26.6, PostgreSQL 18, Apache Kafka 4.3.1"],
    ], widths=[1.6, 5.4])
    page_break(doc)
    para(doc, "Table of Contents", bold=True, size=16, color=NAVY)
    add_toc(doc)
    page_break(doc)
    for b in C:
        k = b[0]
        if k in ("h1", "h2", "h3"):
            if k == "h1" and b[1][0] in "4567" and b[1][1] == ".":
                page_break(doc)
            h(doc, b[1], int(k[1]))
        elif k == "p":
            para(doc, b[1])
        elif k == "note":
            p = para(doc, "Note: " + b[1], italic=True, size=9.5, color=GREY)
            p.paragraph_format.left_indent = Inches(0.2)
        elif k == "bullets":
            bullets(doc, b[1])
        elif k == "code":
            code(doc, b[2], size=8)
        elif k == "table":
            n = len(b[1])
            widths = {2: [2.4, 4.6], 3: [1.9, 2.4, 2.7], 4: [1.6, 1.8, 1.9, 1.7]}.get(n)
            table(doc, b[1], b[2], widths=widths, size=8)
        elif k == "steps":
            for i, (title, text, lang, src) in enumerate(b[1], 1):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(f"Step {i}  ")
                r.bold, r.font.color.rgb, r.font.size = True, NAVY, Pt(11)
                r2 = p.add_run(title)
                r2.bold, r2.font.size = True, Pt(11)
                if text:
                    para(doc, text)
                if src:
                    code(doc, src, size=8.5)
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    OUT_MD.write_text(to_markdown(), encoding="utf-8", newline="\n")
    print(OUT_MD)
    to_docx()
